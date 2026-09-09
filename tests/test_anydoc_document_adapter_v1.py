from __future__ import annotations

import csv
from pathlib import Path

import pytest
from docx import Document as WordDocument
from openpyxl import Workbook
from pptx import Presentation

from backend.document_loader import DocumentLoader
from backend.document_parsers.anydoc_parser import AnyDocParser
from backend.document_parsers.parser_registry import ParserRegistry
from backend.document_parsers.unified_document import to_langchain_documents


def _create_fixture(path: Path) -> None:
    if path.suffix == ".docx":
        document = WordDocument()
        document.add_heading("Annual results", level=1)
        document.add_paragraph("Revenue was 125 million dollars in 2025.")
        document.save(path)
    elif path.suffix == ".xlsx":
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Results"
        sheet.append(["Metric", "2025"])
        sheet.append(["Revenue", 125])
        workbook.save(path)
    elif path.suffix == ".pptx":
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "Annual results"
        slide.placeholders[1].text = "Revenue was 125 million dollars in 2025."
        presentation.save(path)
    elif path.suffix == ".csv":
        with path.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow(["Metric", "2025"])
            writer.writerow(["Revenue", "125"])
    else:  # pragma: no cover - protects the test helper itself
        raise AssertionError(path.suffix)


@pytest.fixture(params=["test.docx", "test.xlsx", "test.pptx", "test.csv"])
def anydoc_fixture(tmp_path: Path, request: pytest.FixtureRequest) -> Path:
    path = tmp_path / str(request.param)
    _create_fixture(path)
    return path


def test_anydoc_fixtures_parse_to_langchain_documents(anydoc_fixture: Path):
    parsed = AnyDocParser().parse(anydoc_fixture)
    documents = to_langchain_documents(parsed)

    assert parsed.parser_backend == "firecrawl-anydoc"
    assert documents
    assert any("125" in document.page_content for document in documents)
    for document in documents:
        assert document.metadata["parser_backend"] == "firecrawl-anydoc"
        assert document.metadata["location"]
        assert "section" in document.metadata
        assert document.metadata["source"] == str(anydoc_fixture)


def test_registry_routes_only_v1_anydoc_extensions():
    registry = ParserRegistry([AnyDocParser(converter=lambda _path: "# Section\nBody")])

    for filename in ("test.docx", "test.xlsx", "test.pptx", "test.csv"):
        assert registry.parser_for(filename) is not None
    for filename in ("test.pdf", "test.txt", "test.md", "legacy.doc", "legacy.xls"):
        assert registry.parser_for(filename) is None


def test_anydoc_documents_continue_through_existing_chunk_pipeline(anydoc_fixture: Path):
    loader = DocumentLoader(chunk_size=64, chunk_overlap=16, include_parent_chunks=False)
    bundle = loader.load_document_bundle(str(anydoc_fixture), anydoc_fixture.name)

    assert bundle["chunks"]
    assert bundle["pages"]
    assert bundle["tables"] == []
    for chunk in bundle["chunks"]:
        assert chunk["parser_backend"] == "firecrawl-anydoc"
        assert chunk["location"]
        assert "section" in chunk
        assert chunk["document_id"]
        assert chunk["page_id"]
        assert chunk["chunk_id"]


def test_txt_stays_on_legacy_loader(tmp_path: Path):
    path = tmp_path / "test.txt"
    path.write_text("Existing text ingestion remains unchanged.", encoding="utf-8")

    document_type, loader = DocumentLoader._resolve_doc_type_and_loader(str(path), path.name)

    assert document_type == "Text"
    assert loader.__class__.__name__ == "TextLoader"

